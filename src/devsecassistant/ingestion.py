
from pathlib import Path
from typing import Any

from docx import Document
from openpyxl import load_workbook
from pptx import Presentation
from pypdf import PdfReader


SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".pptx", ".xlsx"}


def extract_document(path: str | Path) -> list[dict[str, Any]]:
    """Extract text while preserving source-location metadata."""
    path = Path(path)
    extension = path.suffix.lower()

    if not path.is_file():
        raise FileNotFoundError(f"Document not found: {path}")

    if extension not in SUPPORTED_EXTENSIONS:
        raise ValueError(f"Unsupported document type: {extension}")

    if extension == ".pdf":
        return _extract_pdf(path)
    if extension == ".docx":
        return _extract_docx(path)
    if extension == ".pptx":
        return _extract_pptx(path)
    if extension == ".xlsx":
        return _extract_xlsx(path)

    raise ValueError(f"Unsupported document type: {extension}")


def _record(
    path: Path,
    location_type: str,
    location: str,
    text: str,
) -> dict[str, Any]:
    """Create a consistent record for each source location."""
    return {
        "source_file": path.name,
        "source_path": str(path),
        "file_type": path.suffix.lower().lstrip("."),
        "location_type": location_type,
        "location": location,
        "text": text.strip(),
    }


def _extract_pdf(path: Path) -> list[dict[str, Any]]:
    reader = PdfReader(str(path))
    records = []

    for page_number, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        if text.strip():
            records.append(
                _record(path, "page", str(page_number), text)
            )

    return records


def _extract_docx(path: Path) -> list[dict[str, Any]]:
    document = Document(str(path))
    text_parts = [
        paragraph.text
        for paragraph in document.paragraphs
        if paragraph.text.strip()
    ]

    for table in document.tables:
        for row in table.rows:
            values = [
                cell.text.strip()
                for cell in row.cells
            ]
            if any(values):
                text_parts.append(" | ".join(values))

    text = "\n".join(text_parts)
    if not text.strip():
        return []

    return [_record(path, "document", "body", text)]


def _extract_pptx(path: Path) -> list[dict[str, Any]]:
    presentation = Presentation(str(path))
    records = []

    for slide_number, slide in enumerate(
        presentation.slides, start=1
    ):
        parts = []

        for shape in slide.shapes:
            if shape.has_text_frame and shape.text.strip():
                parts.append(shape.text.strip())

            if shape.has_table:
                for row in shape.table.rows:
                    parts.append(
                        " | ".join(cell.text.strip() for cell in row.cells)
                    )

        text = "\n".join(parts)
        if text.strip():
            records.append(
                _record(path, "slide", str(slide_number), text)
            )

    return records


def _extract_xlsx(path: Path) -> list[dict[str, Any]]:
    workbook = load_workbook(
        filename=path,
        read_only=True,
        data_only=False,
    )
    records = []

    try:
        for sheet in workbook.worksheets:
            rows = []

            for row in sheet.iter_rows(values_only=True):
                values = [
                    str(value) if value is not None else ""
                    for value in row
                ]
                if any(value.strip() for value in values):
                    rows.append(" | ".join(values))

            text = "\n".join(rows)
            if text.strip():
                records.append(
                    _record(
                        path,
                        "sheet",
                        sheet.title,
                        text,
                    )
                )
    finally:
        workbook.close()

    return records


def ingest_directory(
    directory: str | Path,
) -> list[dict[str, Any]]:
    """Extract all supported documents in a directory."""
    directory = Path(directory)

    if not directory.is_dir():
        raise NotADirectoryError(
            f"Directory not found: {directory}"
        )

    records = []

    for path in sorted(directory.rglob("*")):
        if not path.is_file():
            continue

        if path.is_symlink():
            continue

        if path.suffix.lower() not in SUPPORTED_EXTENSIONS:
            continue

        records.extend(extract_document(path))

    return records
